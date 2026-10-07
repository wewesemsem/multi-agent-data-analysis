"""Tests for Drafting Agent PDF / DOCX export presentation layer."""

from __future__ import annotations

import io
import sys
import uuid
from pathlib import Path

import pytest

ROOT = Path(__file__).resolve().parent.parent
sys.path.insert(0, str(ROOT))

from app.state import SharedWorkspace
from app.tools import export_tools


def _sample_draft(**overrides):
    draft = {
        "id": f"draft_{uuid.uuid4().hex[:8]}",
        "type": "financial_report",
        "title": "Company A Financial Report",
        "content": (
            "# Company A Financial Report\n\n"
            "## Executive Summary\n\n"
            "Revenue reached **104.5** in 2024.\n\n"
            "- Strong historical growth\n"
            "- Monitor leverage closely\n\n"
            "1. Review forecast assumptions\n"
            "2. Investigate anomalies\n\n"
            "| Year | Revenue |\n"
            "| --- | ---: |\n"
            "| 2023 | 98.2 |\n"
            "| 2024 | 104.5 |\n"
        ),
        "sections": [
            {
                "heading": "Executive Summary",
                "kind": "interpretation",
                "body": (
                    "Revenue reached **104.5** in 2024.\n\n"
                    "- Strong historical growth\n"
                    "- Monitor leverage closely"
                ),
            },
            {
                "heading": "Key Findings",
                "kind": "evidence",
                "body": (
                    "| Year | Revenue |\n"
                    "| --- | ---: |\n"
                    "| 2023 | 98.2 |\n"
                    "| 2024 | 104.5 |\n"
                ),
            },
            {
                "heading": "Next Steps",
                "kind": "recommendation",
                "body": "1. Review forecast assumptions\n2. Investigate anomalies",
            },
        ],
        "tables": [
            {
                "headers": ["Metric", "Value"],
                "rows": [["EBITDA", "26.1"], ["Debt", "32.0"]],
                "caption": "Selected metrics",
            }
        ],
        "charts": ["chart_demo"],
        "source_artifacts": [
            {"kind": "analysis", "index": 0},
            {"kind": "forecast", "index": 0},
            {"kind": "anomaly", "index": 0},
            {"kind": "visualization", "index": 0, "id": "chart_demo"},
        ],
        "assumptions": ["Holdout window used for method selection"],
        "warnings": ["Short history may reduce forecast reliability"],
        "limitations": ["No formal credit score methodology was applied"],
        "grounded": True,
        "evidence_sufficient": True,
        "metadata": {},
    }
    draft.update(overrides)
    return draft


def _workspace_with_chart() -> SharedWorkspace:
    ws = SharedWorkspace()
    ws.analysis_results = [
        {
            "question": "Financial Performance Analysis",
            "query_plan": {"mode": "aggregation"},
            "result": {"records": [{"Year": 2024, "company_revenue": 104.5}], "grounded": True},
            "explanation": "Revenue 104.5",
            "grounded": True,
        }
    ]
    ws.forecasts = [
        {
            "id": "fc1",
            "suitable": True,
            "target_column": "company_revenue",
            "selected_method": "trend_regression",
            "forecast_values": [{"time": "2025", "value": 110.0}],
        }
    ]
    ws.anomalies = [
        {
            "method": "iqr",
            "column": "company_revenue",
            "n_anomalies": 2,
            "grounded": True,
        }
    ]
    ws.visualizations = [
        {
            "id": "chart_demo",
            "chart_type": "line",
            "title": "Revenue Forecast Chart",
            "x": "Year",
            "y": "company_revenue",
            "n_points": 4,
            "grounded": True,
            "data_preview": [
                {"Year": "2021", "company_revenue": 80},
                {"Year": "2022", "company_revenue": 90},
                {"Year": "2023", "company_revenue": 98.2},
                {"Year": "2024", "company_revenue": 104.5},
            ],
        }
    ]
    return ws


def _pdf_text(pdf_bytes: bytes) -> str:
    # reportlab embeds text as literal strings; extract readable ASCII chunks
    raw = pdf_bytes.decode("latin-1", errors="ignore")
    # Also try pypdf if available; otherwise use simple extraction
    try:
        from pypdf import PdfReader

        reader = PdfReader(io.BytesIO(pdf_bytes))
        return "\n".join((page.extract_text() or "") for page in reader.pages)
    except Exception:
        pass
    # Fallback: pull long parenthetical text tokens from content streams
    parts = re_findall_strings(raw)
    return "\n".join(parts)


def re_findall_strings(raw: str) -> list[str]:
    import re

    # PDF literal strings (...)
    found = re.findall(r"\((?:\\.|[^\\)]){3,}\)", raw)
    cleaned = []
    for item in found:
        s = item[1:-1]
        s = s.replace("\\(", "(").replace("\\)", ")").replace("\\n", "\n")
        cleaned.append(s)
    return cleaned


def _docx_text(docx_bytes: bytes) -> str:
    from docx import Document

    document = Document(io.BytesIO(docx_bytes))
    chunks = [p.text for p in document.paragraphs if p.text]
    for table in document.tables:
        for row in table.rows:
            chunks.append(" | ".join(c.text for c in row.cells))
    return "\n".join(chunks)


def _docx_table_count(docx_bytes: bytes) -> int:
    from docx import Document

    return len(Document(io.BytesIO(docx_bytes)).tables)


def _docx_image_count(docx_bytes: bytes) -> int:
    from docx import Document

    document = Document(io.BytesIO(docx_bytes))
    # inline shapes / related image parts
    return len(document.inline_shapes)


# ----- PDF -----


def test_pdf_generated_successfully():
    draft = _sample_draft()
    ws = _workspace_with_chart()
    pdf = export_tools.export_draft_pdf(draft, ws)
    assert isinstance(pdf, bytes)
    assert pdf[:4] == b"%PDF"
    assert len(pdf) > 500


def test_pdf_contains_title_headings_and_text():
    draft = _sample_draft()
    pdf = export_tools.export_draft_pdf(draft, _workspace_with_chart())
    text = _pdf_text(pdf)
    assert "Company A Financial Report" in text
    assert "Executive Summary" in text
    assert "104.5" in text
    assert "Assumptions" in text or "Holdout window" in text


def test_pdf_includes_tables():
    draft = _sample_draft()
    pdf = export_tools.export_draft_pdf(draft, _workspace_with_chart())
    text = _pdf_text(pdf)
    assert "EBITDA" in text
    assert "26.1" in text
    assert "Revenue" in text or "98.2" in text


def test_pdf_includes_charts_when_available():
    draft = _sample_draft()
    ws = _workspace_with_chart()
    doc = export_tools.build_export_document(draft, ws)
    assert any(b.kind == "image" and b.image and b.image.png_bytes[:8] == b"\x89PNG\r\n\x1a\n" for b in doc.blocks)
    pdf = export_tools.export_draft_pdf(draft, ws)
    # PNG image objects are embedded in the PDF stream
    assert b"PNG" in pdf or b"/Image" in pdf


def test_pdf_preserves_assumptions_warnings_limitations():
    draft = _sample_draft()
    text = _pdf_text(export_tools.export_draft_pdf(draft, _workspace_with_chart()))
    assert "Holdout window used for method selection" in text
    assert "Short history may reduce forecast reliability" in text
    assert "No formal credit score methodology was applied" in text


# ----- DOCX -----


def test_docx_generated_successfully():
    draft = _sample_draft()
    docx = export_tools.export_draft_docx(draft, _workspace_with_chart())
    assert isinstance(docx, bytes)
    assert docx[:2] == b"PK"  # zip/docx container
    assert len(docx) > 1000


def test_docx_contains_title_headings_and_text():
    draft = _sample_draft()
    text = _docx_text(export_tools.export_draft_docx(draft, _workspace_with_chart()))
    assert "Company A Financial Report" in text
    assert "Executive Summary" in text
    assert "104.5" in text


def test_docx_includes_tables():
    draft = _sample_draft()
    docx = export_tools.export_draft_docx(draft, _workspace_with_chart())
    assert _docx_table_count(docx) >= 1
    text = _docx_text(docx)
    assert "EBITDA" in text
    assert "26.1" in text


def test_docx_includes_charts_when_available():
    draft = _sample_draft()
    ws = _workspace_with_chart()
    docx = export_tools.export_draft_docx(draft, ws)
    assert _docx_image_count(docx) >= 1


def test_docx_preserves_assumptions_warnings_limitations():
    draft = _sample_draft()
    text = _docx_text(export_tools.export_draft_docx(draft, _workspace_with_chart()))
    assert "Holdout window used for method selection" in text
    assert "Short history may reduce forecast reliability" in text
    assert "No formal credit score methodology was applied" in text
    assert "Sources" in text
    assert "Financial Performance Analysis" in text


# ----- Edge cases -----


def test_export_without_charts():
    draft = _sample_draft(charts=[], source_artifacts=[{"kind": "analysis", "index": 0}])
    ws = SharedWorkspace()
    ws.analysis_results = _workspace_with_chart().analysis_results
    pdf = export_tools.export_draft_pdf(draft, ws)
    docx = export_tools.export_draft_docx(draft, ws)
    assert pdf[:4] == b"%PDF"
    assert docx[:2] == b"PK"
    doc = export_tools.build_export_document(draft, ws)
    assert not any(b.kind == "image" for b in doc.blocks)


def test_export_without_tables():
    draft = _sample_draft(
        tables=[],
        sections=[
            {
                "heading": "Summary",
                "kind": "interpretation",
                "body": "No tabular data in this short note.",
            }
        ],
        content="## Summary\n\nNo tabular data in this short note.",
    )
    pdf = export_tools.export_draft_pdf(draft, SharedWorkspace())
    docx = export_tools.export_draft_docx(draft, SharedWorkspace())
    assert pdf[:4] == b"%PDF"
    assert _docx_table_count(docx) == 0


def test_export_markdown_lists():
    draft = _sample_draft(
        sections=[
            {
                "heading": "Actions",
                "kind": "recommendation",
                "body": "- Alpha item\n- Beta item\n\n1. First\n2. Second",
            }
        ]
    )
    text = _docx_text(export_tools.export_draft_docx(draft, SharedWorkspace()))
    assert "Alpha item" in text
    assert "Beta item" in text
    assert "First" in text
    assert "Second" in text


def test_export_long_content_and_special_characters():
    long_body = "Paragraph with special chars: café—naïve • <tag> & co. " * 80
    draft = _sample_draft(
        title='Q1 Review «Acme / "Beta"» — 100%',
        sections=[{"heading": "Long Section", "kind": "evidence", "body": long_body}],
        content=long_body,
        tables=[],
        charts=[],
    )
    pdf = export_tools.export_draft_pdf(draft, SharedWorkspace())
    docx = export_tools.export_draft_docx(draft, SharedWorkspace())
    assert pdf[:4] == b"%PDF"
    assert "Long Section" in _docx_text(docx)
    fname = export_tools.safe_export_filename(draft["title"], "pdf")
    assert fname.endswith(".pdf")
    assert " " not in fname
    assert "/" not in fname
    assert '"' not in fname


def test_export_failure_does_not_remove_draft():
    ws = _workspace_with_chart()
    draft = _sample_draft()
    ws.drafts = [draft]

    # Force a hard failure inside export and ensure the displayed draft remains.
    original = export_tools._render_pdf

    def _fail(_doc):
        raise RuntimeError("pdf renderer exploded")

    export_tools._render_pdf = _fail  # type: ignore[assignment]
    try:
        with pytest.raises(RuntimeError, match="pdf renderer exploded"):
            export_tools.export_draft_pdf(draft, ws)
    finally:
        export_tools._render_pdf = original  # type: ignore[assignment]

    assert ws.drafts
    assert ws.drafts[0]["id"] == draft["id"]
    assert ws.drafts[0]["content"] == draft["content"]


def test_safe_export_filename():
    assert export_tools.safe_export_filename("My Report!!", "pdf") == "My_Report.pdf"
    assert export_tools.safe_export_filename("", "docx") == "draft.docx"
    long = "A" * 200
    assert len(export_tools.safe_export_filename(long, "pdf")) <= 84
